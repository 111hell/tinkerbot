# /// script
# requires-python = ">=3.13"
# dependencies = ["markitdown[docx,pdf,pptx,xlsx]==0.1.7"]
# ///
"""Generate disposable proof documents for the opt-in Go integration test."""

import sys
from pathlib import Path
from zipfile import ZipFile

from openpyxl import Workbook
from pptx import Presentation

root = Path(sys.argv[1])
root.mkdir(parents=True, exist_ok=True)
proof = "Document extraction proof"
(root / "proof.txt").write_text(proof + "\n中文正文", encoding="utf-8")
(root / "proof.csv").write_text("title,value\n" + proof + ",42\n", encoding="utf-8")

with ZipFile(root / "proof.docx", "w") as archive:
    archive.writestr("[Content_Types].xml", '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>')
    archive.writestr("_rels/.rels", '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>')
    archive.writestr("word/document.xml", '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>' + proof + '</w:t></w:r></w:p></w:body></w:document>')

slides = Presentation()
slide = slides.slides.add_slide(slides.slide_layouts[1])
slide.shapes.title.text = proof
slides.save(root / "proof.pptx")

book = Workbook()
book.active.append(["title", "value"])
book.active.append([proof, 42])
book.save(root / "proof.xlsx")


def pdf(path: Path, text: str) -> None:
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        f"<< /Length {len(stream)} >>\nstream\n".encode() + stream + b"\nendstream",
    ]
    data = bytearray(b"%PDF-1.4\n")
    offsets = []
    for index, body in enumerate(objects, 1):
        offsets.append(len(data))
        data.extend(f"{index} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = len(data)
    data.extend(b"xref\n0 6\n0000000000 65535 f \n")
    for offset in offsets:
        data.extend(f"{offset:010d} 00000 n \n".encode())
    data.extend(f"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    path.write_bytes(data)


pdf(root / "proof.pdf", proof)
pdf(root / "empty.pdf", "")
